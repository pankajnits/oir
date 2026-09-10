#!/usr/bin/env python3
"""Measurement adapters (CEO graph, CSV/Excel tables) — not the product API.

Patterns borrowed:
  Text-to-SQL / MAC-SQL  → table schema → binder
  PyRAG                 → compile plan → execute
  Vault / KG loaders    → triples

You supply: how rows become triples/cells + which compiler family.
"""

from __future__ import annotations

import csv
import json
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from . import EntitySeal, Program, SealRouter, Triple, path_program


@dataclass
class GraphDataset:
    """Generic multi-hop graph: list of records + edge builder."""

    records: list[dict[str, Any]]
    # each record yields triples via edge_fn
    edge_fn: Any
    id_key: str = "id"
    meta: dict = field(default_factory=dict)

    def triples_for(self, record: dict) -> list[Triple]:
        return list(self.edge_fn(record))

    def lexicon(self, field: str) -> set[str]:
        return {str(r[field]) for r in self.records if field in r}

    def sealed_context(self, record: dict, sealer: EntitySeal) -> tuple[list[Triple], str]:
        triples = [sealer.triple(*e) for e in self.triples_for(record)]
        return triples, SealRouter(triples).render()


def load_json_records(path: Path | str) -> list[dict]:
    data = json.loads(Path(path).read_text())
    if isinstance(data, list):
        return data
    if isinstance(data, dict) and "records" in data:
        return data["records"]
    raise ValueError(f"unsupported JSON shape: {path}")


def ceo_hq_edges(row: dict) -> list[Triple]:
    """Default family used in pilots — one of many possible edge_fns."""
    p, c, hq = row["person"], row["company"], row["hq"]
    dco, dhq = f"DecoyCo_{c}", f"DecoyHQ_{hq}"
    return [
        (p, "works_at", c),
        (c, "headquartered_in", hq),
        (dco, "headquartered_in", dhq),
        (c, "owned_by", f"Hold_{c}"),
        (f"Hold_{c}", "meta_of", c),
        (c, "partner_of", f"Partner_{c}"),
        (f"Partner_{c}", "meta_of", c),
        (f"Decoy_{dco}", "works_at", dco),
    ]


def load_ceo_hq_graph(path: Path | str) -> GraphDataset:
    recs = load_json_records(path)
    return GraphDataset(recs, ceo_hq_edges, id_key="person", meta={"family": "ceo_hq"})


@dataclass
class TableDataset:
    """CSV / rectangular table → row-col triples (WTQ-style)."""

    header: list[str]
    rows: list[list[str]]
    meta: dict = field(default_factory=dict)

    def triples(self, max_rows: int | None = None) -> list[Triple]:
        edges: list[Triple] = []
        limit = len(self.rows) if max_rows is None else min(max_rows, len(self.rows))
        for ri in range(limit):
            row_id = f"ROW_{ri + 1}"
            for ci, cell in enumerate(self.rows[ri]):
                if ci >= len(self.header):
                    break
                val = str(cell).strip()
                if not val or val.lower() in ("none", "null", "-"):
                    continue
                col = EntitySeal.normalize(self.header[ci]) or f"COL_{ci}"
                edges.append((row_id, f"col_{col}", EntitySeal.normalize(val)))
        return edges

    def find_answer_cells(self, target: str) -> list[tuple[str, str, str]]:
        """Return (row_id, col_rel, val_atom) hits."""
        tgt = EntitySeal.normalize(target.split("|")[0].strip())
        hits = []
        for h, r, t in self.triples():
            if t == tgt:
                hits.append((h, r, t))
        return hits


def load_csv_table(path: Path | str) -> TableDataset:
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        rows = list(csv.reader(f))
    if not rows:
        return TableDataset([], [])
    return TableDataset(rows[0], rows[1:], meta={"path": str(path)})


def load_excel_table(path: Path | str, sheet: str | int = 0) -> TableDataset:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError(
            "openpyxl required for Excel adapter. Run: pip install 'oir-layer[excel]'"
        ) from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet]
    grid = [[("" if c.value is None else str(c.value)) for c in row] for row in ws.iter_rows()]
    wb.close()
    if not grid:
        return TableDataset([], [], meta={"path": str(path)})
    return TableDataset(grid[0], grid[1:], meta={"path": str(path), "sheet": sheet})


def _col_index(header: list[str], needle: str) -> int:
    needle = needle.lower()
    for i, h in enumerate(header):
        if needle in h.lower():
            return i
    raise KeyError(f"column needle={needle!r} not in {header}")


def excel_rows_as_records(
    path: Path | str,
    *,
    sheet: str | int = 0,
    field_map: dict[str, str] | None = None,
) -> list[dict[str, str]]:
    """Map a single Excel sheet → record dicts.

    field_map: logical_name → header substring (case-insensitive).
    """
    table = load_excel_table(path, sheet)
    fmap = field_map or {
        "person": "employee",
        "company": "company",
    }
    idx = {k: _col_index(table.header, needle) for k, needle in fmap.items()}
    out = []
    for row in table.rows:
        rec = {k: EntitySeal.normalize(row[i]) if i < len(row) else "" for k, i in idx.items()}
        if all(rec.values()):
            out.append(rec)
    return out


def excel_join_employee_company(
    path: Path | str,
    *,
    emp_sheet: str | int = "Employees",
    co_sheet: str | int = "Companies",
    emp_map: dict[str, str] | None = None,
    co_map: dict[str, str] | None = None,
) -> list[dict[str, str]]:
    """Join Employees ⋊ Companies on company name → person/company/hq records.

    Generic 2-sheet HR pattern used by vault-style Excel dumps.
    """
    emp = load_excel_table(path, emp_sheet)
    co = load_excel_table(path, co_sheet)
    em = emp_map or {"person": "employee", "company": "company"}
    cm = co_map or {"company": "company", "hq": "hq"}
    # HQ column may be named HQ / City / Headquarters
    if "hq" in cm:
        for alt in ("hq", "city", "headquarter", "location"):
            try:
                _col_index(co.header, alt)
                cm = {**cm, "hq": alt}
                break
            except KeyError:
                continue
    e_person = _col_index(emp.header, em["person"])
    e_co = _col_index(emp.header, em["company"])
    c_co = _col_index(co.header, cm["company"])
    c_hq = _col_index(co.header, cm["hq"])
    hq_by_co = {}
    for row in co.rows:
        if c_co < len(row) and c_hq < len(row):
            hq_by_co[EntitySeal.normalize(row[c_co])] = EntitySeal.normalize(row[c_hq])
    out = []
    for row in emp.rows:
        if e_person >= len(row) or e_co >= len(row):
            continue
        person = EntitySeal.normalize(row[e_person])
        company = EntitySeal.normalize(row[e_co])
        hq = hq_by_co.get(company, "")
        if person and company and hq:
            out.append({"person": person, "company": company, "hq": hq})
    return out


def make_path_items(
    graph: GraphDataset,
    records: Sequence[dict],
    *,
    start_field: str,
    path: Sequence[str],
    expect_field: str,
    sealer: EntitySeal,
    form: str = "PROG",
) -> list[tuple[str, Program, str, str]]:
    """Build (id, program, sealed_ctx, expect_seal) for MUT or SealRouter."""
    items = []
    for i, rec in enumerate(records):
        start = str(rec[start_field])
        expect_plain = str(rec[expect_field])
        prog = path_program(start, path, expect_plain)
        sealed_prog = prog.seal(sealer)
        # also seal start/rels inside meta for router
        sealed_prog.meta["start"] = sealer.atom(start)
        sealed_prog.meta["rels"] = [sealer.atom(r) for r in path]
        triples, ctx = graph.sealed_context(rec, sealer)
        expect = sealer.atom(expect_plain)
        sealed_prog.expect = expect
        items.append((f"{form}_{i}", sealed_prog, ctx, expect))
    return items
